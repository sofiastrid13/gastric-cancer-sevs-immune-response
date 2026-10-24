#!/usr/bin/env python3
"""Create two clinical tables from TOTAL_TABLE.xlsx.

Install: python -m pip install numpy pandas scipy openpyxl
Run without arguments to select input and output in native file dialogs.
Optional CLI: python create_clinical_tables.py input.xlsx -o results.xlsx
All statistics are recalculated; screenshot numbers are never hardcoded.
"""
from pathlib import Path
import argparse
import itertools
import math
import re
import sys
import warnings

import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SEED = 20261024
BOOTSTRAP_SAMPLES = 10000
PERMUTATIONS = 99999
MIN_INFERENCE_N = 8
PROTEIN = 'Highest_concentration (µg/mL)'
PARTICLES = 'Particles/mL'
HLA = 'sHLA-G_level (ng/mL)'
REQUIRED = ['Patient_ID', 'Sex', 'Age', 'Malignant_pathology_type',
            'Helicobacter_pylori', 'Differentiation_grade',
            'Lymph_node_involvement', 'Tumor_location',
            'Metastatic_involvement', 'TNM', 'Stage', HLA,
            'Vital_status', PROTEIN, PARTICLES]


def load_data(path):
    sheets = pd.read_excel(path, sheet_name=None, keep_default_na=False)
    matches = []
    for name, frame in sheets.items():
        frame.columns = [str(c).strip().replace('μ', 'µ') for c in frame.columns]
        if set(REQUIRED).issubset(frame.columns):
            matches.append((name, frame))
    if len(matches) != 1:
        raise ValueError('Expected exactly one worksheet with these columns:\n' + ', '.join(REQUIRED))
    name, d = matches[0]
    d = d.loc[d.Patient_ID.astype(str).str.strip().ne('')].copy()
    if d.empty or d.Patient_ID.duplicated().any():
        raise ValueError('Patient IDs must be nonempty and unique.')
    for col in ['Age', PROTEIN, PARTICLES, HLA]:
        raw = d[col].astype(str).str.strip()
        values = pd.to_numeric(d[col], errors='coerce')
        bad = values.isna() & ~raw.str.upper().isin(['', 'NA', 'ND', 'NAN', 'NO_DATA', 'NO DATA'])
        if bad.any():
            raise ValueError(f'Invalid numeric values in {col}: {raw[bad].tolist()}')
        if np.isinf(values).any() or (values.dropna() < 0).any():
            raise ValueError(f'Invalid negative or infinite values in {col}.')
        d[col] = values
    known = ['Healthy', 'Chronic_gastritis', 'Chronic_gastritis_(metaplasia)',
             'Peptic_ulcer', 'Stage_I', 'Stage_II', 'Stage_III', 'Stage_IV']
    if not d.Stage.isin(known).all():
        raise ValueError('Unrecognized Stage values: ' + str(d.loc[~d.Stage.isin(known), 'Stage'].tolist()))
    d['Group'] = np.where(d.Stage.eq('Healthy'), 'HC', np.where(d.Stage.str.startswith('Stage_'), 'GC', 'CG'))
    for component in ['T', 'N']:
        d[component] = d.TNM.astype(str).str.upper().str.extract(component + r'([0-4])', expand=False)
    return d, name


def select(d, col, value):
    return d.loc[d[col].eq(value)]


def table1_sections(d):
    gc = select(d, 'Group', 'GC')
    result = []
    for label, mask in [('≤ 60', d.Age.le(60)), ('> 60', d.Age.gt(60))]:
        result.append(('Age (years)', [(f'{label} {g}', d.loc[mask & d.Group.eq(g)]) for g in ['HC','CG','GC']], 'KW'))
    for label, code in [('Female','F'),('Male','M')]:
        result.append(('Sex', [(f'{label} {g}', d.loc[d.Sex.eq(code) & d.Group.eq(g)]) for g in ['HC','CG','GC']], 'KW'))
    result += [
        ('Clinical group', [('CG', d.loc[d.Stage.isin(['Chronic_gastritis','Peptic_ulcer'])]),
                            ('CGm', select(d,'Stage','Chronic_gastritis_(metaplasia)'))], 'MW'),
        ('Clinical stage', [('HC',select(d,'Group','HC'))] + [('GC '+s,select(d,'Stage','Stage_'+s)) for s in ['I','II','III','IV']], 'KW'),
        ('Histological subtype', [('GC (histological diagnosis not recorded)',gc)], None),
        ("Lauren's subtype", [(v,select(gc,'Malignant_pathology_type',v)) for v in ['Diffuse','Intestinal','Mixed']], 'KW'),
        ('Nodal involvement', [(v,select(gc,'Lymph_node_involvement',v)) for v in ['Yes','No','No_data']], 'MW'),
        ('Vital status', [(v,select(gc,'Vital_status',v)) for v in ['Deceased','Alive']], 'MW')]
    return result


def table2_sections(d):
    gc = select(d,'Group','GC')
    groups = [('Clinical group', [('Whole cohort',d)] + [(g,select(d,'Group',g)) for g in ['HC','CG','GC']]),
        ('Clinical characteristics', [(label,select(d,'Stage',v)) for label,v in [
            ('CG without metaplasia','Chronic_gastritis'),('CG with metaplasia','Chronic_gastritis_(metaplasia)'),('Peptic ulcer','Peptic_ulcer')]]),
        ('Clinical stage', [('GC '+s,select(d,'Stage','Stage_'+s)) for s in ['I','II','III','IV']])]
    groups += [(title,rows) for title,rows,_ in table1_sections(d)[:4]]
    hp = gc.Helicobacter_pylori.astype(str).str.lower().str.replace('_',' ')
    cg = select(d,'Group','CG')
    hp_cg = cg.Helicobacter_pylori.astype(str).str.lower().str.replace('_',' ')
    groups += [('Helicobacter pylori', [('Negative CG',cg.loc[hp_cg.isin(['no','absent'])]),
        ('Negative GC',gc.loc[hp.isin(['no','absent'])]),('Positive GC',gc.loc[hp.eq('yes')]),
        ('Unknown CG',cg.loc[~hp_cg.isin(['no','absent','yes'])]),('Unknown GC',gc.loc[~hp.isin(['no','absent','yes'])])])]
    for title,col,values in [
        ("Lauren's subtype",'Malignant_pathology_type',['Diffuse','Intestinal','Mixed']),
        ('Histological differentiation','Differentiation_grade',['Well_differentiated','Moderately_differentiated','Poorly_differentiated']),
        ('Tumor location','Tumor_location',['Gastroesophageal_junction','Cardia','Body','Body_and_antrum','Cardia_and_body','Antrum_and_pylorus']),
        ('Tumor (T)','T',['1','2','3','4']),('Lymph nodes (N)','N',['0','1','2','3']),
        ('Nodal involvement','Lymph_node_involvement',['No','Yes']),
        ('Metastatic involvement','Metastatic_involvement',['No','Yes']),
        ('Vital status','Vital_status',['Alive','Deceased'])]:
        groups.append((title,[((col+v if col in ['T','N'] else v.replace('_',' ')),select(gc,col,v)) for v in values]))
    return groups


def omnibus(rows, col, method):
    if method is None:
        return None
    # Unknown nodal status is descriptive only.
    arrays = [f[col].dropna().to_numpy() for label,f in rows if label != 'No_data']
    arrays = [x for x in arrays if len(x)]
    if len(arrays) < 2:
        return None
    if len(np.unique(np.concatenate(arrays))) == 1:
        return 1.0
    if method == 'MW':
        return float(stats.mannwhitneyu(*arrays, alternative='two-sided', method='auto').pvalue)
    return float(stats.kruskal(*arrays).pvalue)


def correlation(frame):
    pairs = frame[[PROTEIN,HLA]].dropna().to_numpy(float)
    n = len(pairs)
    if n < 2 or any(len(np.unique(pairs[:,i])) < 2 for i in [0,1]):
        return [n,'NE','NR','NR','Not estimable',0]
    x,y = pairs.T
    rho = float(stats.spearmanr(x,y).statistic)
    if n < MIN_INFERENCE_N:
        return [n,rho,'NR','NR','Descriptive only (n < 8)',0]
    rng = np.random.default_rng(SEED)
    ix = rng.integers(0,n,(BOOTSTRAP_SAMPLES,n))
    a,b = stats.rankdata(x[ix],axis=1),stats.rankdata(y[ix],axis=1)
    a -= a.mean(axis=1,keepdims=True)
    b -= b.mean(axis=1,keepdims=True)
    den = np.sqrt((a*a).sum(1)*(b*b).sum(1))
    boot = (a*b).sum(1)[den>0]/den[den>0]
    lo,hi = np.quantile(boot,[.025,.975])
    # Permute ranks of one variable. This preserves ties and uses |rho|
    # for a two-sided test. Exact enumeration for n <= 9.
    rx,ry = stats.rankdata(x),stats.rankdata(y)
    rx -= rx.mean(); ry -= ry.mean()
    denominator = np.linalg.norm(rx)*np.linalg.norm(ry)
    threshold = abs(rho)-1e-12
    extreme = 0
    if n <= 9:
        iterator = itertools.permutations(ry)
        while True:
            batch = list(itertools.islice(iterator,10000))
            if not batch: break
            extreme += int(np.count_nonzero(np.abs(np.asarray(batch)@rx/denominator)>=threshold))
        p = extreme/math.factorial(n)
        method = 'Exact permutation (two-sided)'
    else:
        for start in range(0,PERMUTATIONS,5000):
            count = min(5000,PERMUTATIONS-start)
            perms = rng.permuted(np.broadcast_to(ry,(count,n)),axis=1)
            extreme += int(np.count_nonzero(np.abs(perms@rx/denominator)>=threshold))
        p = (extreme+1)/(PERMUTATIONS+1)
        method = 'Monte Carlo permutation (two-sided)'
    return [n,rho,f'({lo:.3f}, {hi:.3f})',p,method,len(boot)]


def style_sheet(ws, widths, landscape=False):
    ws.sheet_view.showGridLines = False
    for col,width in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws:
        for cell in row:
            cell.font = Font(name='Times New Roman',size=11)
            cell.alignment = Alignment(vertical='center',wrap_text=True)
    ws.freeze_panes = 'B5'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if landscape else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:4'
    ws.print_options.horizontalCentered = True
    ws.print_area = ws.dimensions


def heading(ws,row,text,end):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=end)
    c = ws.cell(row,1,text)
    c.font = Font(name='Times New Roman',size=11,bold=True)
    c.fill = PatternFill('solid',fgColor='E7EDF3')
    ws.row_dimensions[row].height = 23


def build(input_path, output_path):
    if input_path.resolve() == output_path.resolve():
        raise ValueError('The output must be different from the input file.')
    d, source_sheet = load_data(input_path)
    wb = Workbook()
    t1 = wb.active; t1.title = 'Table 1'
    t2 = wb.create_sheet('Table 2')
    details = wb.create_sheet('Calculation details')
    details.append(['Section','Subgroup','Measure','n','Median','Q1','Q3','Test','p-value'])
    t1.merge_cells('A1:J1'); t1['A1'] = 'Table 1. Clinicopathological characteristics'
    for start,title in [(2,'sEV protein concentration (µg/mL)'),(5,'sEV concentration (particles/mL)'),(8,'sHLA-G concentration (ng/mL)')]:
        t1.merge_cells(start_row=3,start_column=start,end_row=3,end_column=start+2)
        t1.cell(3,start,title)
    t1.append(['Characteristic','n','Median (Q1–Q3)','p','n','Median (Q1–Q3)','p','n','Median (Q1–Q3)','p'])
    r=5
    last=None
    for title,rows,method in table1_sections(d):
        if title != last:
            heading(t1,r,title,10); r+=1; last=title
        ps = [omnibus(rows,c,method) for c in [PROTEIN,PARTICLES,HLA]]
        for j,(label,frame) in enumerate(rows):
            t1.cell(r,1,label.replace('No_data','No data'))
            for i,col in enumerate([PROTEIN,PARTICLES,HLA]):
                v=frame[col].dropna().to_numpy(float)
                q = np.quantile(v,[.25,.5,.75]) if len(v) else [None]*3
                fmt = (lambda a:f'{a:.2E}') if col == PARTICLES else (lambda a:f'{a:.3f}')
                summary = f'{fmt(q[1])} ({fmt(q[0])}–{fmt(q[2])})' if len(v) else 'ND'
                p = ps[i] if j==0 else None
                for k,value in enumerate([len(v),summary,p]): t1.cell(r,2+3*i+k,value)
                if p is not None: t1.cell(r,4+3*i).number_format='0.0000'
                details.append([title,label,col,len(v),q[1],q[0],q[2],method,p])
            t1.row_dimensions[r].height=32
            r+=1
    t2.merge_cells('A1:E1'); t2['A1']='Table 2. Spearman correlations: sEV protein vs. sHLA-G'
    t2.merge_cells('A2:E2'); t2['A2']='Paired nonmissing measurements; 95% percentile bootstrap confidence intervals'
    for i,v in enumerate(['Subgroup','n','Spearman rho','95% bootstrap CI','p-value'],1): t2.cell(4,i,v)
    r=5
    corr_audit=[]
    cache={}
    last=None
    for title,rows in table2_sections(d):
        if title!=last: heading(t2,r,title,5); r+=1; last=title
        for label,frame in rows:
            key=tuple(frame.index)
            if key not in cache: cache[key]=correlation(frame)
            n,rho,ci,p,method,valid=cache[key]
            for i,v in enumerate([label,n,rho,ci,p],1): t2.cell(r,i,v)
            t2.cell(r,3).number_format='0.000'; t2.cell(r,5).number_format='0.00000'
            corr_audit.append([title,label,n,method,valid])
            t2.row_dimensions[r].height=24
            r+=1
    notes = wb.create_sheet('Methods')
    lines = [
        'Methods and data interpretation',
        f'Source: {input_path.name}; worksheet: {source_sheet}; patient count: {len(d)}.',
        'Each source row represents one patient. Units are retained exactly as supplied. No concentration scaling is applied.',
        'HC: healthy controls. GC: gastric cancer. CG in age/sex/cohort analyses includes all nonmalignant gastric conditions, including peptic ulcer.',
        'Table 1 CG combines Chronic_gastritis and Peptic_ulcer to follow the reference grouping; CGm denotes gastritis with metaplasia. Table 2 lists peptic ulcer separately.',
        'The source does not explicitly record adenocarcinoma. The histological-subtype row is therefore labeled as diagnosis not recorded rather than asserting adenocarcinoma.',
        'Unknown H. pylori values are not classified as negative. Negative means No or Absent. Unknown groups are shown separately.',
        'Table 1: medians and linear-interpolation Q1/Q3. n counts available values for each measure, not missing measurements.',
        'MW: two-sided Mann–Whitney U, SciPy auto method (exact for a small sample without ties, otherwise asymptotic with tie and continuity correction).',
        'KW: Kruskal–Wallis across nonempty groups. Age and sex comparisons are among clinical groups within each stratum. HC vs GC stages uses five groups.',
        'Each p-value is a block-level comparison, displayed on the first subgroup row. Unknown nodal involvement is excluded from the Yes/No test.',
        'The reference has additional unlabeled p-values in its clinical-group block. Their hypotheses cannot be recovered from the image. No unlabeled comparisons or post-hoc tests are invented.',
        'Table 2: Spearman rho from paired nonmissing values. NE means not estimable (<2 pairs or constant values); NR means not reported.',
        f'For n < {MIN_INFERENCE_N}, rho is descriptive only, following the supplied reference. This cutoff is a reporting convention, not a guarantee of reliable inference.',
        f'For n >= {MIN_INFERENCE_N}: {BOOTSTRAP_SAMPLES:,} paired bootstrap resamples; percentile 95% CI. Constant resamples are excluded, with valid count recorded below.',
        f'Two-sided permutation p-values use absolute rho: exact enumeration for n <= 9; otherwise {PERMUTATIONS:,} random permutations with (extreme+1)/(B+1). Seed: {SEED}.',
        'P-values are unadjusted exploratory results. Values below 0.05 are bold. Monte Carlo and bootstrap values may differ from the screenshot.',
        'TNM components are parsed independently, including T1a and T4b as T1 and T4. Unknown X components are excluded. Source stage and nodal fields are not overwritten from TNM.',
        'Outputs are calculated values. Rerun the script after changing the source data. Calculation details retain unrounded numerical summaries.',
    ]
    for line in lines: notes.append([line])
    notes.append([]); notes.append(['Section','Subgroup','Paired n','Correlation p-value method','Valid bootstrap resamples'])
    for row in corr_audit: notes.append(row)
    style_sheet(t1,[42,7,38,12,7,43,12,7,38,12],True)
    style_sheet(t2,[37,7,16,25,14])
    style_sheet(details,[25,37,37,8,21,21,21,12,16],True)
    style_sheet(notes,[42,38,16,48,26],True)
    details.print_title_rows = '1:1'
    notes.print_title_rows = '1:1'
    for row in details:
        details.row_dimensions[row[0].row].height = 32
    for ws, columns in [(t1,[2,5,8]),(t2,[2,3,5])]:
        for row in ws.iter_rows(min_row=5):
            for col in columns:
                row[col-1].alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
    for row in range(1,len(lines)+1):
        notes.merge_cells(start_row=row,start_column=1,end_row=row,end_column=5)
        notes.row_dimensions[row].height=42
    for ws in [t1,t2]:
        ws.row_dimensions[1].height=32
        for cell in ws[1]+ws[3]+ws[4]: cell.font=Font(name='Times New Roman',size=11,bold=True)
        for row in ws:
            for c in row:
                if isinstance(c.value,(float,int)) and c.column in ([4,7,10] if ws==t1 else [5]) and c.row>=5 and c.value<.05:
                    c.font=Font(name='Times New Roman',size=11,bold=True)
        # Restore section styling after the general style pass.
        for merged in ws.merged_cells.ranges:
            if merged.min_row>=5:
                ws.cell(merged.min_row,1).font=Font(name='Times New Roman',size=11,bold=True)
    output_path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(output_path)
    return len(d)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input',nargs='?',type=Path)
    parser.add_argument('-o','--output',type=Path)
    args=parser.parse_args()
    root=None
    try:
        if args.input is None:
            import tkinter as tk
            from tkinter import filedialog, messagebox
            root=tk.Tk(); root.withdraw(); root.attributes('-topmost',True)
            name=filedialog.askopenfilename(parent=root,title='Select the TOTAL TABLE input file',filetypes=[('Excel workbook','*.xlsx')])
            if not name: return 0
            args.input=Path(name)
            target=filedialog.asksaveasfilename(parent=root,title='Save the clinical tables',initialdir=str(args.input.parent),initialfile='Clinical_Tables.xlsx',defaultextension='.xlsx',filetypes=[('Excel workbook','*.xlsx')])
            if not target: return 0
            args.output=Path(target)
        if args.output is None: args.output=args.input.with_name(args.input.stem+'_Clinical_Tables.xlsx')
        print('Calculating tables. Please wait...',flush=True)
        count=build(args.input,args.output)
        print(f'Completed: {count} patients. Saved to {args.output.resolve()}')
        if root is not None: messagebox.showinfo('Completed',f'Tables saved successfully:\n{args.output.resolve()}',parent=root)
        return 0
    except Exception as error:
        text=f'{type(error).__name__}: {error}\nClose the output workbook in Excel before saving. Verify the required columns and installed packages.'
        print(text,file=sys.stderr)
        if root is not None:
            from tkinter import messagebox
            messagebox.showerror('Unable to create tables',text,parent=root)
        return 1
    finally:
        if root is not None: root.destroy()


if __name__=='__main__':
    sys.exit(main())
