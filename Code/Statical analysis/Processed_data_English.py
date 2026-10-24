#!/usr/bin/env python3
"""Create processed checkpoint and cytokine tables from concentration data.

EASY START (Windows / VS Code):
    Run this file using Run Python File, then select your concentration Excel file.
    No command-line arguments or fixed file names are required.
    Output is saved beside the selected workbook. Existing outputs are preserved.

OPTIONAL COMMAND-LINE USAGE:
    python Processed_data_English.py concentrations.xlsx
    python Processed_data_English.py concentrations.xlsx -o results.xlsx
    python Processed_data_English.py concentrations.xlsx --reference original.xlsx

DEPENDENCIES: numpy, pandas, openpyxl, scipy. Missing packages are reported with
an installation command for the exact interpreter running this script.

METHODS:
    Technical replicate count is not a count of independent donors.
    Summary: median +/- SEM, where SEM = sample SD / sqrt(n); SEM is ND for n < 2.
    Q1/Q3 use linear interpolation, equivalent to Excel PERCENTILE.INC.
    FC = group mean / HC mean; log2FC = log2(FC).
    Percentage change = 100 * (FC - 1).
    Absolute difference = group mean - HC mean (signed, matching the source table).
    Symmetric difference = (group mean - HC mean) / (group mean + HC mean).
    Undefined results are ND; missing measurements are excluded, never zero-filled.
    The original SD heading is corrected to SEM. Correlations are recalculated
    using matching replicate labels and may differ from the original workbook.
    Replicate labels alone do not establish biological pairing.
"""
from pathlib import Path
import argparse
import math
import warnings
import sys
import subprocess
import importlib.util


def load_dependencies():
    required = ['numpy', 'pandas', 'openpyxl', 'scipy']
    missing = [name for name in required if importlib.util.find_spec(name) is None]
    if missing:
        command = subprocess.list2cmdline([sys.executable, '-m', 'pip', 'install', *missing])
        print('Missing packages: ' + ', '.join(missing))
        print('Install them using this exact Python interpreter:')
        print(('& ' if sys.platform == 'win32' else '') + command)
        print('Then run this script again.')
        return False
    return True


if not load_dependencies():
    sys.exit(1)

import numpy as np
import pandas as pd
from scipy.stats import spearmanr
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

GROUPS = ['HC', 'CG', 'GC I', 'GC II', 'GC III', 'GC IV', 'SC', 'UC']
FIELDS = ['n', 'median ± SEM', 'Q1 - Q3', 'FC vs HC', 'log2FC',
          'Percentage change', 'Absolute difference', 'Symmetric difference']

def number(x):
    if x is None or str(x).strip().upper() in {'', 'ND', 'NA', 'NAN'}:
        return np.nan
    try:
        return float(str(x).strip().replace(',', '.'))
    except ValueError:
        raise ValueError(f'Unrecognized concentration: {x!r}')

def read_source(path, name, first_marker):
    df = pd.read_excel(path, sheet_name=name, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    df = df.loc[df.iloc[:, 0].notna()].copy()
    df.iloc[:, 0] = df.iloc[:, 0].astype(str).str.strip()
    df['Replicate'] = df['Replicate'].astype(str).str.strip()
    df['Time'] = df['Time'].map(number)
    markers = list(df.columns[first_marker:])
    for col in markers:
        df[col] = df[col].map(number)
    keys = list(df.columns[:first_marker])
    if df.duplicated(keys).any():
        raise ValueError(f'Duplicate rows in {name}; check group, replicate, time and stimulation.')
    return df, markers

def select(df, group, time, stimulation=None):
    mask = (df.iloc[:, 0] == group) & (df['Time'] == time)
    if stimulation is not None:
        mask &= df['Stimulation'].astype(str).str.strip().str.upper().eq(stimulation)
    return df.loc[mask]

def stats(values, ref):
    a = np.asarray(values, dtype=float)
    a = a[np.isfinite(a)]
    b = np.asarray(ref, dtype=float)
    b = b[np.isfinite(b)]
    if not len(a):
        return [0] + ['ND'] * 7
    n = len(a)
    median = np.median(a)
    sem = np.std(a, ddof=1) / np.sqrt(n) if n > 1 else np.nan
    q1, q3 = np.quantile(a, [.25, .75], method='linear')
    summary = f'{median:.3f} ± {sem:.3f}' if n > 1 else f'{median:.3f} ± ND'
    result = [n, summary, f'{q1:.3f} - {q3:.3f}']
    if not len(b):
        return result + ['ND'] * 5
    x, h = np.mean(a), np.mean(b)
    fc = x / h if h != 0 else np.nan
    log = math.log2(fc) if fc > 0 else np.nan
    symmetric = (x-h)/(x+h) if x+h != 0 else np.nan
    return result + [float(v) if np.isfinite(v) else 'ND'
                     for v in [fc, log, (fc-1)*100, x-h, symmetric]]

def correlation(df, group, time, marker):
    a = select(df, group, time)[['Replicate', marker]]
    b = select(df, 'sEVs ' + group, 0)[['Replicate', marker]]
    pairs = a.merge(b, on='Replicate', suffixes=('_pbmc', '_sev')).dropna()
    n = len(pairs)
    if n < 2 or pairs.iloc[:, 1].nunique() < 2 or pairs.iloc[:, 2].nunique() < 2:
        return ['ND', n]
    return [float(spearmanr(pairs.iloc[:, 1], pairs.iloc[:, 2]).statistic), n]

def build(wb, df, markers, checkpoints):
    ws = wb.create_sheet('Checkpoints' if checkpoints else 'Cytokines')
    conditions = [('24 h',24,None), ('48 h',48,None), ('72 h',72,None),
                  ('sEVs',0,None)] if checkpoints else [
                      ('6 h',6,'U'), ('36 h (SC)',36,'S'),
                      ('36 h (UC)',36,'U'), ('72 h',72,'U')]
    ws.cell(1,1,'Marker / clinical group')
    for j,(label,_,_) in enumerate(conditions):
        start = 2+j*9
        ws.cell(1,start,label + ' (pg/mL)')
        ws.merge_cells(start_row=1,start_column=start,end_row=1,end_column=start+7)
        for k,f in enumerate(FIELDS): ws.cell(2,start+k,f)
    if checkpoints:
        for j,t in enumerate([24,48,72]):
            ws.cell(1,39+j*3,f'PBMCs vs sEVs: {t} h')
            ws.cell(2,39+j*3,'Spearman rho')
            ws.cell(2,40+j*3,'Matched technical pairs')
    row = 3
    for marker in markers:
        ws.cell(row,1,'B7.2' if marker=='B7.2 (CD86)' else marker)
        row += 1
        for group in GROUPS:
            ws.cell(row,1,group)
            for j,(_,time,stim) in enumerate(conditions):
                sev = checkpoints and j==3
                g = 'sEVs '+group if sev else group
                h = 'sEVs HC' if sev else 'HC'
                vals = stats(select(df,g,time,stim)[marker], select(df,h,time,stim)[marker])
                for k,v in enumerate(vals): ws.cell(row,2+j*9+k,v)
            if checkpoints:
                for j,time in enumerate([24,48,72]):
                    for k,v in enumerate(correlation(df,group,time,marker)):
                        ws.cell(row,39+j*3+k,v)
            row += 1
    ws.freeze_panes='B3'
    ws.column_dimensions['A'].width=24
    for col in range(2,ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width=22
    for cells in ws.iter_rows():
        for c in cells:
            c.alignment=Alignment(vertical='center')
            if isinstance(c.value,float): c.number_format='0.000000'
            if c.row <=2 or (c.column==1 and c.value not in GROUPS):
                c.fill=PatternFill('solid',fgColor='17365D')
                c.font=Font(color='FFFFFF',bold=True)
    return ws

def compare(output, reference):
    """Compare the five change metrics, ignoring formatting and external formulas."""
    old = load_workbook(reference, data_only=True)
    new = load_workbook(output, data_only=True)
    checked=0; differences=[]
    for name in ['Checkpoints','Cytokines']:
        original=old[name]; generated=new[name]
        offset=2 if name=='Checkpoints' else 1
        for r in range(4,generated.max_row+1):
            if generated.cell(r,1).value not in GROUPS: continue
            for block in range(4):
                for k in range(5):
                    col=5+block*9+k
                    a=original.cell(r+offset,col).value
                    b=generated.cell(r,col).value
                    if isinstance(a,(float,int)) and isinstance(b,(float,int)):
                        checked+=1
                        if not np.isclose(a,b,rtol=1e-8,atol=1e-8):
                            differences.append(f'{name}!{get_column_letter(col)}{r+offset}: original={a}; calculated={b}')
    report=Path(output).with_suffix('.validation.txt')
    report.write_text(f'Numeric change metrics compared: {checked}\nDifferences: {len(differences)}\n'+'\n'.join(differences),encoding='utf-8')
    print(f'Validation: {checked} metrics; {len(differences)} differences. {report}')

def choose_input():
    """Open a file picker; use a typed path if Tk is unavailable."""
    root = None
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        selected = filedialog.askopenfilename(
            title='Select the concentration Excel workbook',
            initialdir=str(Path(__file__).resolve().parent),
            filetypes=[('Excel workbooks', '*.xlsx *.xlsm')])
        return Path(selected) if selected else None
    except Exception:
        print('The file picker is unavailable.')
        try:
            value = input('Paste the full concentration workbook path (Enter to cancel): ').strip().strip('"').strip("'")
            return Path(value) if value else None
        except EOFError:
            return None
    finally:
        if root is not None:
            root.destroy()


def unique_output(path):
    """Keep previous output files instead of overwriting them."""
    if not path.exists():
        return path
    index = 2
    while path.with_name(f'{path.stem}_{index}{path.suffix}').exists():
        index += 1
    return path.with_name(f'{path.stem}_{index}{path.suffix}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', nargs='?', type=Path,
                        help='Optional concentration workbook. Omit to open a file picker.')
    parser.add_argument('-o', '--output', type=Path,
                        help='Optional output path. Existing files are preserved.')
    parser.add_argument('--reference', type=Path,
                        help='Optional original workbook for validation only.')
    args = parser.parse_args()
    source = args.input.expanduser() if args.input else choose_input()
    if source is None:
        print('No workbook selected. Nothing was changed.')
        return 0
    if not source.is_file():
        print(f'Workbook not found: {source}')
        print('Please select the workbook in the file picker.')
        source = choose_input()
        if source is None:
            print('Cancelled. Nothing was changed.')
            return 0
    if not source.is_file():
        print(f'Workbook not found: {source}')
        return 1
    source = source.resolve()
    destination = args.output or source.parent / 'Processed_Data_recalculated.xlsx'
    destination = destination.expanduser().resolve()
    if destination == source:
        print('The output must be different from the concentration workbook.')
        return 1
    if destination.suffix.lower() != '.xlsx':
        print('The output filename must end with .xlsx.')
        return 1
    destination = unique_output(destination)
    try:
        cp, cm = read_source(source, 'Checkpoints', 3)
        cy, ym = read_source(source, 'Cytokines', 4)
        wb = Workbook()
        wb.remove(wb.active)
        build(wb, cp, cm, True)
        build(wb, cy, ym, False)
        destination.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destination)
    except PermissionError:
        print('Access denied. Close the workbook in Excel and check folder permissions, then try again.')
        return 1
    except (ValueError, KeyError) as error:
        print(f'Workbook structure or data error: {error}')
        print('Select the concentration workbook containing the Checkpoints and Cytokines sheets.')
        return 1
    except Exception as error:
        print(f'Unable to process the workbook ({type(error).__name__}): {error}')
        return 1
    print(f'Workbook created successfully:\n{destination}')
    print('Sheets: Checkpoints and Cytokines.')
    print('Errors are reported as SEM. Fold changes use means and HC as reference.')
    print('ND values are excluded, never replaced with zero. No significance tests are performed.')
    print('Spearman correlations use matching replicate labels; confirm experimental pairing before interpretation.')
    if args.reference:
        try:
            compare(destination, args.reference)
        except Exception as error:
            print(f'The output was saved, but optional validation failed: {error}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
