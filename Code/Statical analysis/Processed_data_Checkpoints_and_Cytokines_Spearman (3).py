#!/usr/bin/env python3
"""Create the processed checkpoint and cytokine tables (concentrations and
Spearman correlations) from the concentration workbook.

EASY START (Windows / VS Code):
    Run this file using Run Python File, then select your concentration Excel file.
    No command-line arguments or fixed file names are required.
    Output is saved beside the selected workbook. Existing outputs are preserved.

OPTIONAL COMMAND-LINE USAGE:
    python Processed_data_Checkpoints_and_Cytokines_Spearman.py concentrations.xlsx
    python ... concentrations.xlsx -o results.xlsx
    python ... concentrations.xlsx --reference original.xlsx
    python ... concentrations.xlsx --n-boot 10000 --n-perm 10000 --seed 2026

DEPENDENCIES: numpy, pandas, openpyxl, scipy. Missing packages are reported with
an installation command for the exact interpreter running this script.

OUTPUT SHEETS (same tables as Processed_Data):
    Checkpoints            Concentrations (24/48/72 h, sEVs) + PBMCs vs sEVs Spearman.
    Spearman checkpoints   Checkpoint vs checkpoint, per group, 24/48/72 h.
    Cytokines              Concentrations (6 h, 36 h SC, 36 h UC, 72 h).
    Spearman Cytokines     Cytokine vs cytokine, per group, 6 / 36 UC / 36 SC / 72 h.
    Spearman 72 h          Checkpoints + cytokines at 72 h, per group.
    Methods                Parameters.

CONCENTRATION METHODS:
    Summary: median +/- SEM, where SEM = sample SD / sqrt(n); SEM is ND for n < 2.
    Q1/Q3 use linear interpolation (Excel PERCENTILE.INC).
    FC = group mean / HC mean; log2FC = log2(FC); Percentage change = 100 * (FC - 1).
    Absolute difference = group mean - HC mean.
    Symmetric difference = (group mean - HC mean) / (group mean + HC mean).
    Missing measurements are excluded, never zero-filled. Undefined results are ND.

CORRELATION METHODS:
    Every correlation reports rho, n, Monte Carlo permutation p and 95% bootstrap CI.
    Spearman rank correlation (average ranks for ties); pairwise deletion of
    missing values; rows paired by group + replicate (+ time / stimulation).
    p: two-sided Monte Carlo permutation test, p = (b + 1) / (B + 1).
    95% CI: percentile bootstrap, resampling pairs with replacement; resamples in
    which a variable is constant (rho undefined) are discarded.
    Significance: p < 0.05 (highlighted in yellow, bold).
    rho is ND when n < 2 or one variable is constant.
    Replicate labels alone do not establish biological pairing.
"""
from pathlib import Path
import argparse
import itertools
import math
import sys
import subprocess
import importlib.util
import zlib


def load_dependencies():
    required = ['numpy', 'pandas', 'openpyxl', 'scipy']
    missing = [name for name in required if importlib.util.find_spec(name) is None]
    if missing:
        command = subprocess.list2cmdline([sys.executable, '-m', 'pip', 'install', *missing])
        print('Missing packages: ' + ', '.join(missing))
        print('Install them using this exact Python interpreter:')
        print(('& ' if sys.platform == 'win32' else '') + command)
        print('Then run this script again.')
        return False
    return True


if not load_dependencies():
    sys.exit(1)

import numpy as np
import pandas as pd
from scipy.stats import spearmanr, rankdata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# Settings
# ----------------------------------------------------------------------------
GROUPS = ['HC', 'CG', 'GC I', 'GC II', 'GC III', 'GC IV', 'SC', 'UC']
FIELDS = ['n', 'median ± SEM', 'Q1 - Q3', 'FC vs HC', 'log2FC',
          'Percentage change', 'Absolute difference', 'Symmetric difference']
CORR_FIELDS = ['ρ', 'n', 'p (Monte Carlo)', '95% bootstrap CI']

N_BOOT = 10000
N_PERM = 10000
ALPHA = 0.05
SEED = 20261025
TOL = 1e-9

HEADER_FILL = PatternFill('solid', fgColor='17365D')
SUB_FILL = PatternFill('solid', fgColor='DCE6F1')
HEADER_FONT = Font(color='FFFFFF', bold=True)
SIG_FILL = PatternFill('solid', fgColor='FFE699')
SIG_FONT = Font(bold=True, color='9C0006')
THIN = Side(style='thin', color='BFBFBF')


# ----------------------------------------------------------------------------
# Reading
# ----------------------------------------------------------------------------
def number(x):
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return np.nan
    if str(x).strip().upper() in {'', 'ND', 'NA', 'NAN'}:
        return np.nan
    try:
        return float(str(x).strip().replace(',', '.'))
    except ValueError:
        raise ValueError(f'Unrecognized concentration: {x!r}')


def read_source(path, name, first_marker):
    df = pd.read_excel(path, sheet_name=name, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    df = df.loc[df.iloc[:, 0].notna()].copy()
    df.iloc[:, 0] = df.iloc[:, 0].astype(str).str.strip()
    df['Replicate'] = df['Replicate'].astype(str).str.strip()
    df['Time'] = df['Time'].map(number)
    if 'Stimulation' in df.columns:
        df['Stimulation'] = df['Stimulation'].astype(str).str.strip().str.upper()
    markers = list(df.columns[first_marker:])
    for col in markers:
        df[col] = df[col].map(number).astype(float)
    keys = list(df.columns[:first_marker])
    if df.duplicated(keys).any():
        raise ValueError(f'Duplicate rows in {name}; check group, replicate, time and stimulation.')
    return df, markers


def combined_72h(path, cp, cp_markers, cy, cy_markers):
    """Checkpoints + cytokines at 72 h. Uses the 'C & C 72 h' sheet when present
    and reports any disagreement with the Checkpoints/Cytokines sheets;
    otherwise merges those two sheets."""
    notes = []
    if 'C & C 72 h' in pd.ExcelFile(path).sheet_names:
        df, markers = read_source(path, 'C & C 72 h', 3)
        a = cp[cp['Time'] == 72].set_index([cp.columns[0], 'Replicate'])
        b = cy[(cy['Time'] == 72) & (cy['Stimulation'] == 'U')].set_index([cy.columns[0], 'Replicate'])
        d = df.set_index([df.columns[0], 'Replicate'])
        for m in markers:
            src = a if m in a.columns else (b if m in b.columns else None)
            if src is None:
                continue
            for idx in d.index:
                x = d.at[idx, m]
                y = src.at[idx, m] if idx in src.index else np.nan
                if not ((np.isnan(x) and np.isnan(y)) or
                        (np.isfinite(x) and np.isfinite(y) and abs(x - y) < 1e-6)):
                    notes.append(f'C & C 72 h differs from source sheet: {idx[0]} {idx[1]} {m}: '
                                 f'{x if np.isfinite(x) else "ND"} vs {y if np.isfinite(y) else "ND"}')
        return df, markers, "'C & C 72 h' sheet", notes
    a = cp[cp['Time'] == 72][[cp.columns[0], 'Replicate'] + cp_markers]
    b = cy[(cy['Time'] == 72) & (cy['Stimulation'] == 'U')][[cy.columns[0], 'Replicate'] + cy_markers]
    b = b.rename(columns={cy.columns[0]: cp.columns[0]})
    df = a.merge(b, on=[cp.columns[0], 'Replicate'], how='outer')
    df['Time'] = 72.0
    return df, cp_markers + cy_markers, 'Checkpoints (72 h) merged with Cytokines (72 h, U)', notes


def select(df, group, time, stimulation=None):
    mask = (df.iloc[:, 0] == group) & (df['Time'] == time)
    if stimulation is not None:
        mask &= df['Stimulation'].eq(stimulation)
    return df.loc[mask]


# ----------------------------------------------------------------------------
# Concentration statistics
# ----------------------------------------------------------------------------
def stats(values, ref):
    a = np.asarray(values, dtype=float)
    a = a[np.isfinite(a)]
    b = np.asarray(ref, dtype=float)
    b = b[np.isfinite(b)]
    if not len(a):
        return [0] + ['ND'] * 7
    n = len(a)
    median = np.median(a)
    sem = np.std(a, ddof=1) / np.sqrt(n) if n > 1 else np.nan
    q1, q3 = np.quantile(a, [.25, .75], method='linear')
    summary = f'{median:.3f} ± {sem:.3f}' if n > 1 else f'{median:.3f} ± ND'
    result = [n, summary, f'{q1:.3f} - {q3:.3f}']
    if not len(b):
        return result + ['ND'] * 5
    x, h = np.mean(a), np.mean(b)
    fc = x / h if h != 0 else np.nan
    log = math.log2(fc) if fc > 0 else np.nan
    symmetric = (x - h) / (x + h) if x + h != 0 else np.nan
    return result + [float(v) if np.isfinite(v) else 'ND'
                     for v in [fc, log, (fc - 1) * 100, x - h, symmetric]]


# ----------------------------------------------------------------------------
# Spearman: rho, Monte Carlo permutation p, bootstrap CI
# ----------------------------------------------------------------------------
def _rowwise_pearson(a, b):
    a = a - a.mean(axis=1, keepdims=True)
    b = b - b.mean(axis=1, keepdims=True)
    den = np.sqrt((a * a).sum(axis=1) * (b * b).sum(axis=1))
    with np.errstate(invalid='ignore', divide='ignore'):
        return np.where(den > 0, (a * b).sum(axis=1) / den, np.nan)


def spearman_full(x, y, label, cfg):
    """rho, n, two-sided Monte Carlo permutation p and 95% percentile bootstrap CI."""
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    ok = np.isfinite(x) & np.isfinite(y)
    x, y = x[ok], y[ok]
    n = int(len(x))
    out = {'n': n, 'rho': 'ND', 'p': 'ND', 'ci': 'ND', 'sig': False}
    if n < 2 or np.unique(x).size < 2 or np.unique(y).size < 2:
        return out
    out['rho'] = round(float(spearmanr(x, y).statistic), 10)
    rng = np.random.default_rng([cfg['seed'], zlib.crc32(label.encode('utf-8'))])

    # Monte Carlo permutation test (two-sided) on ranks
    rx, ry = rankdata(x), rankdata(y)
    obs = abs(out['rho'])
    B = cfg['n_perm']
    perm_y = rng.permuted(np.tile(ry, (B, 1)), axis=1)
    null = np.abs(_rowwise_pearson(np.broadcast_to(rx, perm_y.shape), perm_y))
    out['p'] = float((np.sum(null >= obs - TOL) + 1) / (B + 1))
    out['sig'] = out['p'] < ALPHA

    # Percentile bootstrap CI (resampling pairs)
    idx = rng.integers(0, n, size=(cfg['n_boot'], n))
    boot = _rowwise_pearson(rankdata(x[idx], axis=1), rankdata(y[idx], axis=1))
    boot = boot[np.isfinite(boot)]
    if boot.size:
        lo, hi = np.percentile(boot, [2.5, 97.5])
        out['ci'] = f'{lo:.3f} to {hi:.3f}'
    return out


def corr_cells(res):
    return [res['rho'], res['n'], res['p'], res['ci']]


# ----------------------------------------------------------------------------
# Writing helpers
# ----------------------------------------------------------------------------
def style_header(cell, fill=HEADER_FILL, font=HEADER_FONT):
    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def write_corr(ws, row, col, res):
    for k, v in enumerate(corr_cells(res)):
        c = ws.cell(row, col + k, v)
        if isinstance(v, float):
            c.number_format = '0.000' if k == 0 else '0.0000'
        c.alignment = Alignment(horizontal='center', vertical='center')
        if res['sig']:
            c.fill = SIG_FILL
            c.font = SIG_FONT


# ----------------------------------------------------------------------------
# Sheet: Checkpoints / Cytokines (concentrations; PBMCs vs sEVs Spearman)
# ----------------------------------------------------------------------------
def build(wb, df, markers, checkpoints, cfg, counter):
    ws = wb.create_sheet('Checkpoints' if checkpoints else 'Cytokines')
    conditions = [('24 h', 24, None), ('48 h', 48, None), ('72 h', 72, None),
                  ('sEVs', 0, None)] if checkpoints else [
        ('6 h', 6, 'U'), ('36 h (SC)', 36, 'S'),
        ('36 h (UC)', 36, 'U'), ('72 h', 72, 'U')]
    ws.cell(1, 1, 'Marker / clinical group')
    for j, (label, _, _) in enumerate(conditions):
        start = 2 + j * 9
        ws.cell(1, start, label + ' (pg/mL)')
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + 7)
        for k, f in enumerate(FIELDS):
            ws.cell(2, start + k, f)
    corr_start = 2 + len(conditions) * 9          # column after the last block + gap
    width = len(CORR_FIELDS) + 1
    if checkpoints:
        for j, t in enumerate([24, 48, 72]):
            c0 = corr_start + j * width
            ws.cell(1, c0, f'PBMCs vs sEVs (Spearman): {t} h')
            ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + len(CORR_FIELDS) - 1)
            for k, f in enumerate(CORR_FIELDS):
                ws.cell(2, c0 + k, f)
    row = 3
    for marker in markers:
        ws.cell(row, 1, 'B7.2' if marker == 'B7.2 (CD86)' else marker)
        row += 1
        for group in GROUPS:
            ws.cell(row, 1, group)
            for j, (_, time, stim) in enumerate(conditions):
                sev = checkpoints and j == 3
                g = 'sEVs ' + group if sev else group
                h = 'sEVs HC' if sev else 'HC'
                vals = stats(select(df, g, time, stim)[marker], select(df, h, time, stim)[marker])
                for k, v in enumerate(vals):
                    c = ws.cell(row, 2 + j * 9 + k, v)
                    if isinstance(v, float):
                        c.number_format = '0.000000'
            if checkpoints:
                for j, t in enumerate([24, 48, 72]):
                    a = select(df, group, t)[['Replicate', marker]]
                    b = select(df, 'sEVs ' + group, 0)[['Replicate', marker]]
                    p = a.merge(b, on='Replicate', suffixes=('_pbmc', '_sev'))
                    res = spearman_full(p[marker + '_pbmc'], p[marker + '_sev'],
                                        f'PBMC-sEV|{t}|{group}|{marker}', cfg)
                    counter.append(res)
                    write_corr(ws, row, corr_start + j * width, res)
            row += 1
    ws.freeze_panes = 'B3'
    ws.column_dimensions['A'].width = 24
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    for cells in ws.iter_rows():
        for c in cells:
            if c.row <= 2 or (c.column == 1 and c.value is not None and c.value not in GROUPS):
                style_header(c)
    return ws


# ----------------------------------------------------------------------------
# Sheets: Spearman checkpoints / Spearman Cytokines / Spearman 72 h
# ----------------------------------------------------------------------------
def build_spearman(wb, title, df, markers, conditions, cfg, counter):
    """Rows: Marker 1 -> time/condition -> Marker 2 (every pair once).
    Columns: one block (ρ, n, p, CI) per clinical group."""
    ws = wb.create_sheet(title)
    width = len(CORR_FIELDS) + 1
    first = 4
    ws.cell(1, 1, 'Spearman correlations: ρ, n, two-sided Monte Carlo permutation p '
                  f'(B = {cfg["n_perm"]}), 95% percentile bootstrap CI ({cfg["n_boot"]} resamples). '
                  f'Highlighted: p < {ALPHA}.')
    ws.cell(1, 1).font = Font(italic=True)
    for h, label in enumerate(['Marker 1', 'Time', 'Marker 2']):
        style_header(ws.cell(2, h + 1, label))
        ws.merge_cells(start_row=2, start_column=h + 1, end_row=3, end_column=h + 1)
    for gi, g in enumerate(GROUPS):
        c0 = first + gi * width
        style_header(ws.cell(2, c0, g))
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + len(CORR_FIELDS) - 1)
        for k, f in enumerate(CORR_FIELDS):
            style_header(ws.cell(3, c0 + k, f), SUB_FILL, Font(bold=True))
    row = 4
    total = len(markers) - 1
    for mi, m1 in enumerate(markers[:-1]):
        style_header(ws.cell(row, 1, m1))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        row += 1
        for cond_label, time, stim in conditions:
            for m2 in markers[mi + 1:]:
                ws.cell(row, 1, m1)
                ws.cell(row, 2, cond_label)
                ws.cell(row, 3, m2)
                for gi, g in enumerate(GROUPS):
                    sub = select(df, g, time, stim)
                    res = spearman_full(sub[m1].values, sub[m2].values,
                                        f'{title}|{cond_label}|{g}|{m1}|{m2}', cfg)
                    counter.append(res)
                    write_corr(ws, row, first + gi * width, res)
                for c in range(1, first + len(GROUPS) * width - 1):
                    ws.cell(row, c).border = Border(bottom=THIN)
                row += 1
        print(f'\r  {title}: {mi + 1}/{total} markers', end='', flush=True)
    print()
    ws.freeze_panes = ws.cell(4, first)
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 8
    ws.column_dimensions['C'].width = 16
    for gi in range(len(GROUPS)):
        c0 = first + gi * width
        for k, w in enumerate([9, 5, 11, 17, 2]):
            ws.column_dimensions[get_column_letter(c0 + k)].width = w
    ws.auto_filter.ref = f'A3:{get_column_letter(first + len(GROUPS) * width - 2)}{row - 1}'
    return ws


def write_methods(wb, cfg, source72, notes, counts):
    ws = wb.create_sheet('Methods')
    lines = [
        'Spearman correlation analysis',
        'All correlations are reported with ρ, n, Monte Carlo permutation p-value and 95% bootstrap CI.',
        'Spearman rank correlation (average ranks for ties); pairwise deletion of missing values.',
        'Rows paired by clinical group + replicate (+ time / stimulation). Replicate labels alone do not establish biological pairing.',
        f'p-value: two-sided Monte Carlo permutation test, B = {cfg["n_perm"]}, p = (b + 1)/(B + 1).',
        f'95% CI: percentile bootstrap, {cfg["n_boot"]} resamples of pairs; resamples with a constant variable are discarded.',
        f'Significance: p < {ALPHA} (highlighted in yellow, bold red text).',
        'ρ = ND when n < 2 or one variable is constant.',
        'Note: with n = 3 only 3! = 6 orderings exist, so the smallest attainable two-sided p is about 0.33;'
        ' with n = 2, ρ is always ±1 and p is about 1.',
        f'Random seed: {cfg["seed"]} (each correlation uses its own derived seed; results are reproducible).',
        f'72 h checkpoint-cytokine data source: {source72}.',
        '',
        'Number of correlations',
    ] + [f'{k}: {v}' for k, v in counts.items()]
    if notes:
        lines += ['', 'Data consistency warnings'] + notes
    for i, line in enumerate(lines, start=1):
        c = ws.cell(i, 1, line)
        if line in ('Spearman correlation analysis', 'Number of correlations', 'Data consistency warnings'):
            c.font = Font(bold=True)
    ws.column_dimensions['A'].width = 140


# ----------------------------------------------------------------------------
# Optional validation against the original workbook (concentration metrics)
# ----------------------------------------------------------------------------
def compare(output, reference):
    old = load_workbook(reference, data_only=True)
    new = load_workbook(output, data_only=True)
    checked = 0
    differences = []
    for name in ['Checkpoints', 'Cytokines']:
        original = old[name]
        generated = new[name]
        offset = 2 if name == 'Checkpoints' else 1
        for r in range(4, generated.max_row + 1):
            if generated.cell(r, 1).value not in GROUPS:
                continue
            for block in range(4):
                for k in range(5):
                    col = 5 + block * 9 + k
                    a = original.cell(r + offset, col).value
                    b = generated.cell(r, col).value
                    if isinstance(a, (float, int)) and isinstance(b, (float, int)):
                        checked += 1
                        if not np.isclose(a, b, rtol=1e-8, atol=1e-8):
                            differences.append(f'{name}!{get_column_letter(col)}{r + offset}: original={a}; calculated={b}')
    report = Path(output).with_suffix('.validation.txt')
    report.write_text(f'Numeric change metrics compared: {checked}\nDifferences: {len(differences)}\n'
                      + '\n'.join(differences), encoding='utf-8')
    print(f'Validation: {checked} metrics; {len(differences)} differences. {report}')


# ----------------------------------------------------------------------------
# File handling
# ----------------------------------------------------------------------------
def choose_input():
    root = None
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        selected = filedialog.askopenfilename(
            title='Select the concentration Excel workbook',
            initialdir=str(Path(__file__).resolve().parent),
            filetypes=[('Excel workbooks', '*.xlsx *.xlsm')])
        return Path(selected) if selected else None
    except Exception:
        print('The file picker is unavailable.')
        try:
            value = input('Paste the full concentration workbook path (Enter to cancel): ').strip().strip('"').strip("'")
            return Path(value) if value else None
        except EOFError:
            return None
    finally:
        if root is not None:
            root.destroy()


def unique_output(path):
    if not path.exists():
        return path
    index = 2
    while path.with_name(f'{path.stem}_{index}{path.suffix}').exists():
        index += 1
    return path.with_name(f'{path.stem}_{index}{path.suffix}')


def summarize(rows):
    valid = [r for r in rows if isinstance(r['rho'], float)]
    sig = sum(r['sig'] for r in valid)
    return f'{len(rows)} correlations; {len(valid)} computed; {sig} with p < {ALPHA}; {len(rows) - len(valid)} ND'


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('input', nargs='?', type=Path,
                        help='Optional concentration workbook. Omit to open a file picker.')
    parser.add_argument('-o', '--output', type=Path,
                        help='Optional output path. Existing files are preserved.')
    parser.add_argument('--reference', type=Path,
                        help='Optional original workbook for validation only.')
    parser.add_argument('--n-boot', type=int, default=N_BOOT, help=f'Bootstrap resamples (default {N_BOOT}).')
    parser.add_argument('--n-perm', type=int, default=N_PERM, help=f'Monte Carlo permutations (default {N_PERM}).')
    parser.add_argument('--seed', type=int, default=SEED, help=f'Random seed (default {SEED}).')
    args = parser.parse_args()
    cfg = {'n_boot': args.n_boot, 'n_perm': args.n_perm, 'seed': args.seed}

    source = args.input.expanduser() if args.input else choose_input()
    if source is None:
        print('No workbook selected. Nothing was changed.')
        return 0
    if not source.is_file():
        print(f'Workbook not found: {source}')
        print('Please select the workbook in the file picker.')
        source = choose_input()
        if source is None:
            print('Cancelled. Nothing was changed.')
            return 0
    if not source.is_file():
        print(f'Workbook not found: {source}')
        return 1
    source = source.resolve()
    destination = args.output or source.parent / 'Processed_Data_recalculated.xlsx'
    destination = destination.expanduser().resolve()
    if destination == source:
        print('The output must be different from the concentration workbook.')
        return 1
    if destination.suffix.lower() != '.xlsx':
        print('The output filename must end with .xlsx.')
        return 1
    destination = unique_output(destination)
    try:
        cp, cm = read_source(source, 'Checkpoints', 3)
        cy, ym = read_source(source, 'Cytokines', 4)
        c72, m72, source72, notes = combined_72h(source, cp, cm, cy, ym)
        print('Computing tables and Spearman correlations...')
        counts = {}
        wb = Workbook()
        wb.remove(wb.active)
        rows = []
        build(wb, cp, cm, True, cfg, rows)
        counts['PBMCs vs sEVs (Checkpoints sheet)'] = summarize(rows)
        rows = []
        build_spearman(wb, 'Spearman checkpoints', cp, cm,
                       [('24 h', 24, None), ('48 h', 48, None), ('72 h', 72, None)], cfg, rows)
        counts['Spearman checkpoints'] = summarize(rows)
        build(wb, cy, ym, False, cfg, [])
        rows = []
        build_spearman(wb, 'Spearman Cytokines', cy, ym,
                       [('6 h', 6, 'U'), ('36 h UC', 36, 'U'), ('36 h SC', 36, 'S'), ('72 h', 72, 'U')],
                       cfg, rows)
        counts['Spearman Cytokines'] = summarize(rows)
        rows = []
        build_spearman(wb, 'Spearman 72 h', c72, m72, [('72 h', 72, None)], cfg, rows)
        counts['Spearman 72 h'] = summarize(rows)
        write_methods(wb, cfg, source72, notes, counts)
        destination.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destination)
    except PermissionError:
        print('Access denied. Close the workbook in Excel and check folder permissions, then try again.')
        return 1
    except (ValueError, KeyError) as error:
        print(f'Workbook structure or data error: {error}')
        print('Select the concentration workbook containing the Checkpoints and Cytokines sheets.')
        return 1
    except Exception as error:
        print(f'Unable to process the workbook ({type(error).__name__}): {error}')
        return 1
    print(f'Workbook created successfully:\n{destination}')
    for k, v in counts.items():
        print(f'  {k}: {v}')
    for n in notes:
        print('  WARNING: ' + n)
    print(f'Significant correlations (p < {ALPHA}) are highlighted in yellow.')
    if args.reference:
        try:
            compare(destination, args.reference)
        except Exception as error:
            print(f'The output was saved, but optional validation failed: {error}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
